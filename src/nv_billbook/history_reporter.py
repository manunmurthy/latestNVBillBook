"""Write standalone multi-month flat collection history workbooks."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
ALT_ROW_FILL = PatternFill("solid", fgColor="F7F9FB")
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="FCE4D6")
BLUE_FILL = PatternFill("solid", fgColor="DDEBF7")
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
HEADER_FONT = Font(bold=True)
THIN_GREY = Side(style="thin", color="D9E1F2")
CELL_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)
MONEY_HEADERS = {
    "Amount",
    "Expected Maintenance",
    "Water Bill",
    "Expected",
    "Paid",
    "Amount Due",
    "Extra Paid",
    "Total Expected",
    "Total Paid",
    "Total Amount Due",
    "Total Extra Paid",
}


def _write_table(
    writer: pd.ExcelWriter,
    sheet_name: str,
    frame: pd.DataFrame,
    start_row: int,
) -> int:
    """Write and style one table, returning the next available row index."""
    frame.to_excel(writer, sheet_name=sheet_name, startrow=start_row, index=False)
    worksheet = writer.sheets[sheet_name]
    header_row = start_row + 1
    for cell in worksheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=header_row + len(frame),
        max_col=len(frame.columns),
    ):
        for cell in row:
            header = worksheet.cell(row=header_row, column=cell.column).value
            if header in MONEY_HEADERS and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            cell.alignment = Alignment(vertical="top", wrap_text=header in {"Payer", "Payment Type"})
            if (cell.row - header_row) % 2 == 1:
                cell.fill = ALT_ROW_FILL

    return start_row + len(frame) + 3


def _set_column_widths(worksheet) -> None:
    for column_cells in worksheet.iter_cols():
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if values:
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(
                max(len(value) for value in values) + 2, 30
            )


def _add_status_colors(worksheet, header_row: int, start_row: int, end_row: int) -> None:
    headers = {
        cell.value: cell.column
        for cell in worksheet[header_row]
        if cell.value is not None
    }
    status_col = headers.get("Status")
    if not status_col:
        return
    letter = get_column_letter(status_col)
    status_range = f"{letter}{start_row}:{letter}{max(start_row, end_row)}"
    for value, fill, color in (
        ("Paid", GREEN_FILL, "008000"),
        ("Excess", BLUE_FILL, "0066CC"),
        ("Pending", YELLOW_FILL, "9C6500"),
        ("Short", RED_FILL, "9C0006"),
    ):
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'{letter}{start_row}="{value}"'],
                fill=fill,
                font=Font(color=color, bold=True),
            ),
        )


def write_collection_history_report(
    output_path: Path,
    transaction_history: pd.DataFrame,
    monthly_reconciliation: pd.DataFrame,
    period_summary: pd.DataFrame,
) -> Path:
    """Write one worksheet per flat with its summary, monthly rows, and payments."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    flat_order = monthly_reconciliation["Flat"].drop_duplicates().tolist()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for flat_no in flat_order:
            sheet_name = str(flat_no)
            flat_summary = period_summary[period_summary["Flat"] == flat_no]
            flat_monthly = monthly_reconciliation[monthly_reconciliation["Flat"] == flat_no]
            flat_transactions = transaction_history[transaction_history["Flat"] == flat_no]

            next_row = _write_table(writer, sheet_name, flat_summary, start_row=2)
            worksheet = writer.sheets[sheet_name]
            worksheet["A1"] = f"Flat {flat_no} — Collection History"
            worksheet["A1"].font = TITLE_FONT
            worksheet["A1"].fill = TITLE_FILL

            worksheet.cell(row=next_row, column=1).value = "Monthly Reconciliation"
            worksheet.cell(row=next_row, column=1).font = HEADER_FONT
            next_row = _write_table(writer, sheet_name, flat_monthly, start_row=next_row)
            monthly_header_row = next_row - len(flat_monthly) - 2
            _add_status_colors(
                worksheet,
                monthly_header_row,
                monthly_header_row + 1,
                next_row - 3,
            )

            worksheet.cell(row=next_row, column=1).value = "Payment Transactions"
            worksheet.cell(row=next_row, column=1).font = HEADER_FONT
            _write_table(writer, sheet_name, flat_transactions, start_row=next_row)

            worksheet.freeze_panes = "A3"
            worksheet.sheet_view.showGridLines = False
            _set_column_widths(worksheet)

    return output_path
