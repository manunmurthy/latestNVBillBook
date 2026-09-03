"""Generate monthly Excel reports from classified transactions."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from nv_billbook.config import Config
from nv_billbook.flats_registry import FlatsRegistry
from nv_billbook.parser import ParsedStatement
from nv_billbook.reconciliation import build_maintenance_reconciliation

MONEY_COLUMNS = ("debit", "credit", "amount", "balance")
SUMMARY_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
ALT_ROW_FILL = PatternFill("solid", fgColor="F7F9FB")
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="FCE4D6")
BLUE_FILL = PatternFill("solid", fgColor="DDEBF7")
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
HEADER_FONT = Font(bold=True)
MUTED_FONT = Font(color="666666")
THIN_GREY = Side(style="thin", color="D9E1F2")
CARD_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)


def _format_sheet_columns(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    worksheet = writer.sheets[sheet_name]
    for idx, column in enumerate(df.columns, start=1):
        values = [str(column)] + [
            str(value)
            for value in df[column].tolist()
            if value is not None and str(value) != "nan"
        ]
        width = min(max(len(value) for value in values) + 2, 60)
        worksheet.column_dimensions[get_column_letter(idx)].width = width

    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
        for row_number in range(2, worksheet.max_row + 1):
            if row_number % 2 == 0:
                for cell in worksheet[row_number]:
                    cell.fill = ALT_ROW_FILL
            for cell in worksheet[row_number]:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 3, 4, 5})

        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = CARD_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                header = worksheet.cell(row=1, column=cell.column).value
                if header in {"Debit (₹)", "Credit (₹)", "Amount (₹)", "Balance (₹)"}:
                    cell.number_format = "#,##0.00"


def _add_summary_cards(worksheet, summary_df: pd.DataFrame) -> None:
    """Add a compact visual snapshot beside the detailed summary table."""
    metrics = {
        str(row["Metric"]): row["Value"]
        for _, row in summary_df.iterrows()
        if str(row["Metric"]).strip()
    }
    cards = [
        ("Total Credits (₹)", "Total Credits", BLUE_FILL),
        ("Total Debits (₹)", "Total Debits", RED_FILL),
        ("Closing Balance (₹)", "Closing Balance", GREEN_FILL),
        ("Pending Flats", "Pending Flats", YELLOW_FILL),
    ]
    for index, (metric_key, label, fill) in enumerate(cards):
        start_col = 4 + index * 2
        end_col = start_col + 1
        start_letter = get_column_letter(start_col)
        end_letter = get_column_letter(end_col)
        worksheet.merge_cells(f"{start_letter}2:{end_letter}2")
        worksheet.merge_cells(f"{start_letter}3:{end_letter}4")
        label_cell = worksheet[f"{start_letter}2"]
        value_cell = worksheet[f"{start_letter}3"]
        label_cell.value = label
        value_cell.value = metrics.get(metric_key, 0)
        label_cell.fill = fill
        value_cell.fill = fill
        label_cell.font = Font(bold=True, color="1F4E79")
        value_cell.font = Font(bold=True, size=14)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        label_cell.border = CARD_BORDER
        value_cell.border = CARD_BORDER
        if metric_key.endswith("(₹)"):
            value_cell.number_format = "#,##0.00"
        worksheet.column_dimensions[start_letter].width = 16
        worksheet.column_dimensions[end_letter].width = 16


def _add_reconciliation_conditional_formatting(worksheet) -> None:
    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    status_col = headers.get("Status")
    flat_col = headers.get("Flat")
    if status_col:
        letter = get_column_letter(status_col)
        status_range = f"{letter}2:{letter}{max(2, worksheet.max_row)}"
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'{letter}2="Paid"'], fill=GREEN_FILL, font=Font(color="008000", bold=True)),
        )
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'{letter}2="Excess"'], fill=BLUE_FILL, font=Font(color="0066CC", bold=True)),
        )
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'{letter}2="Pending"'], fill=YELLOW_FILL, font=Font(color="9C6500", bold=True)),
        )
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(formula=[f'{letter}2="Short"'], fill=RED_FILL, font=Font(color="9C0006", bold=True)),
        )
    if flat_col:
        letter = get_column_letter(flat_col)
        flat_range = f"{letter}2:{letter}{max(2, worksheet.max_row)}"
        worksheet.conditional_formatting.add(
            flat_range,
            FormulaRule(formula=[f'{letter}2=""'], fill=YELLOW_FILL),
        )


def _prepare_display_frame(df: pd.DataFrame, txn_type: str) -> pd.DataFrame:
    if df.empty:
        if txn_type == "credit":
            return pd.DataFrame(
                columns=["Date", "Narration", "Reference", "Flat No", "SBA (sqft)", "Payee", "Credit (₹)", "Balance (₹)"]
            )
        if txn_type == "debit":
            return pd.DataFrame(
                columns=["Date", "Narration", "Reference", "Payee", "Purpose / Type", "Debit (₹)", "Balance (₹)"]
            )
        return pd.DataFrame(
            columns=["Date", "Narration", "Reference", "Payee", "Purpose / Type", "Debit (₹)", "Credit (₹)", "Amount (₹)", "Balance (₹)"]
        )

    def _format_date(row: pd.Series) -> str:
        if isinstance(row["date"], date):
            return row["date"].strftime("%d/%m/%Y")
        return str(row.get("date_str", ""))

    if txn_type == "credit":
        display = pd.DataFrame(
            {
                "Date": df.apply(_format_date, axis=1),
                "Narration": df["narration"],
                "Reference": df["reference"],
                "Flat No": df["flat_no"].fillna(""),
                "SBA (sqft)": df.get("sqft", pd.Series(index=df.index, dtype="object")).apply(
                    lambda v: int(v) if pd.notna(v) and str(v).strip() not in ("", "nan") else ""
                ),
                "Payee": df["counterparty"].fillna(""),
            }
        )
        display["Credit (₹)"] = df["credit"]
        display["Balance (₹)"] = df["balance"]
    elif txn_type == "debit":
        display = pd.DataFrame(
            {
                "Date": df.apply(_format_date, axis=1),
                "Narration": df["narration"],
                "Reference": df["reference"],
                "Payee": df["counterparty"].fillna(""),
                "Purpose / Type": df["purpose"].fillna(""),
            }
        )
        display["Debit (₹)"] = df["debit"]
        display["Balance (₹)"] = df["balance"]
    else:
        display = pd.DataFrame(
            {
                "Date": df.apply(_format_date, axis=1),
                "Narration": df["narration"],
                "Reference": df["reference"],
                "Payee": df["counterparty"].fillna(""),
                "Purpose / Type": df["purpose"].fillna(""),
            }
        )
        display["Debit (₹)"] = df["debit"]
        display["Credit (₹)"] = df["credit"]
        display["Amount (₹)"] = df["amount"]
        display["Balance (₹)"] = df["balance"]

    return display


def _build_summary_rows(
    parsed: ParsedStatement,
    split: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    all_txn = split["all"]
    credits = split["credits"]
    debits = split["debits"]
    other = split["other"]

    opening_balance = None
    if not all_txn.empty and all_txn.iloc[0]["balance"] is not None:
        first = all_txn.iloc[0]
        opening_balance = (first["balance"] or 0) - first["credit"] + first["debit"]

    closing_balance = None
    if not all_txn.empty:
        closing_balance = all_txn.iloc[-1]["balance"]

    period_label = ""
    if parsed.meta.period_start and parsed.meta.period_end:
        period_label = (
            f"{parsed.meta.period_start.strftime('%d/%m/%Y')} to "
            f"{parsed.meta.period_end.strftime('%d/%m/%Y')}"
        )

    rows = [
        ("Report Month", parsed.meta.period_start.strftime("%B %Y") if parsed.meta.period_start else ""),
        ("Statement Period", period_label),
        ("Source File", parsed.meta.source_file.name),
        ("Account No", parsed.meta.account_no or ""),
        ("", ""),
        ("Opening Balance (₹)", opening_balance),
        ("Total Credits (₹)", credits["credit"].sum() if not credits.empty else 0),
        ("Total Debits (₹)", debits["debit"].sum() if not debits.empty else 0),
        ("Other / Charges (₹)", other["debit"].sum() - other["credit"].sum() if not other.empty else 0),
        ("Net Movement (₹)", all_txn["amount"].sum() if not all_txn.empty else 0),
        ("Closing Balance (₹)", closing_balance),
        ("", ""),
        ("Credit Transactions", len(credits)),
        ("Debit Transactions", len(debits)),
        ("Other Transactions", len(other)),
        ("Total Transactions", len(all_txn)),
    ]

    summary = pd.DataFrame(rows, columns=["Metric", "Value"])

    if not debits.empty:
        summary = pd.concat([summary, pd.DataFrame([("", "")], columns=["Metric", "Value"])], ignore_index=True)
        category_totals = (
            debits.groupby("category", dropna=False)["debit"]
            .sum()
            .sort_values(ascending=False)
            .reset_index()
        )
        category_totals.columns = ["Metric", "Value"]
        category_totals["Metric"] = "Expense: " + category_totals["Metric"].astype(str)
        summary = pd.concat([summary, category_totals], ignore_index=True)

    if not credits.empty:
        summary = pd.concat([summary, pd.DataFrame([("", "")], columns=["Metric", "Value"])], ignore_index=True)
        income_totals = (
            credits.groupby("category", dropna=False)["credit"]
            .sum()
            .sort_values(ascending=False)
            .reset_index()
        )
        income_totals.columns = ["Metric", "Value"]
        income_totals["Metric"] = "Income: " + income_totals["Metric"].astype(str)
        summary = pd.concat([summary, income_totals], ignore_index=True)

    return summary


def _style_summary_sheet(writer: pd.ExcelWriter, sheet_name: str) -> None:
    ws = writer.sheets[sheet_name]
    ws["A1"].font = Font(bold=True, size=12)
    ws["B1"].font = Font(bold=True, size=12)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = CARD_BORDER
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20


def write_monthly_report(
    parsed: ParsedStatement,
    split: dict[str, pd.DataFrame],
    config: Config,
    report_month: str,
    flats_registry: FlatsRegistry | None = None,
    water_bill_per_month: object | None = None,
) -> Path:
    """Write a monthly workbook and return the output path."""
    config.output_dir.mkdir(parents=True, exist_ok=True)
    output_path = config.output_dir / f"{report_month}_summary.xlsx"

    summary_df = _build_summary_rows(parsed, split)
    credits_df = _prepare_display_frame(split["credits"], "credit")
    debits_df = _prepare_display_frame(split["debits"], "debit")
    other_df = _prepare_display_frame(split["other"], "other")
    all_df = _prepare_display_frame(split["all"], "all")
    reconciliation_df = (
        build_maintenance_reconciliation(
            split["credits"],
            flats_registry,
            water_bill_per_month=water_bill_per_month,
        )
        if flats_registry
        else pd.DataFrame()
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        credits_df.to_excel(writer, sheet_name="Credits", index=False)
        debits_df.to_excel(writer, sheet_name="Debits", index=False)
        other_df.to_excel(writer, sheet_name="Other", index=False)
        all_df.to_excel(writer, sheet_name="All Transactions", index=False)
        reconciliation_df.to_excel(writer, sheet_name="Maintenance Reconciliation", index=False)

        _style_summary_sheet(writer, "Summary")
        for sheet_name, frame in (
            ("Credits", credits_df),
            ("Debits", debits_df),
            ("Other", other_df),
            ("All Transactions", all_df),
        ):
            _format_sheet_columns(writer, sheet_name, frame)

        _format_sheet_columns(writer, "Maintenance Reconciliation", reconciliation_df)
        reconciliation_sheet = writer.sheets["Maintenance Reconciliation"]
        for cell in reconciliation_sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = CARD_BORDER
        for row in reconciliation_sheet.iter_rows(min_row=2, max_row=reconciliation_sheet.max_row):
            for cell in row:
                if reconciliation_sheet.cell(row=1, column=cell.column).value in {
                    "Expected Maintenance",
                    "Water Bill",
                    "Expected",
                    "Paid",
                    "Amount Due",
                    "Extra Paid",
                }:
                    cell.number_format = "#,##0.00"
        _add_reconciliation_conditional_formatting(reconciliation_sheet)

        summary_sheet = writer.sheets["Summary"]
        summary_sheet.insert_rows(1, 2)
        summary_sheet["A1"] = config.society_name
        summary_sheet["A1"].font = TITLE_FONT
        summary_sheet["A1"].fill = SUMMARY_FILL
        summary_sheet["A2"] = f"Monthly Bank Report — {report_month}"
        summary_sheet["A2"].font = Font(bold=True, size=11)
        summary_sheet.merge_cells("A1:B1")
        summary_sheet.merge_cells("A2:B2")
        _add_summary_cards(summary_sheet, summary_df)

    return output_path
