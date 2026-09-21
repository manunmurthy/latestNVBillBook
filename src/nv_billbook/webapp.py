"""Local browser interface for Nava Vaibhva Bill Book reports."""

from __future__ import annotations

import tempfile
from collections.abc import Mapping
from pathlib import Path

from flask import Flask, abort, render_template_string, request, send_from_directory, url_for
from openpyxl import load_workbook

from nv_billbook.classifier import classify_transactions, split_by_type
from nv_billbook.collection_history import (build_flat_transaction_history, build_monthly_reconciliation, build_period_summary, month_keys_from_transactions)
from nv_billbook.config import Config
from nv_billbook.history import _load_statements, _parse_flat_filter
from nv_billbook.history_reporter import write_collection_history_report
from nv_billbook.main import _load_flats_registry, _process_statement

ALLOWED_EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".csv"}

PAGE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Nava Vaibhva Bill Book</title><style>
:root{--navy:#173b5c;--blue:#e9f2fa;--line:#ccd9e5}*{box-sizing:border-box}body{margin:0;font:16px system-ui,-apple-system,sans-serif;background:#f5f8fb;color:#1d2b38}header{background:var(--navy);color:#fff;padding:28px max(24px,calc((100% - 1050px)/2))}h1{margin:0;font-size:26px}header p{margin:7px 0 0;opacity:.86}main{max-width:1050px;margin:28px auto;padding:0 24px 48px}.notice{padding:14px 16px;border-radius:8px;margin-bottom:20px;background:#e6f5ec;color:#125432}.error{background:#fdeceb;color:#8a251c}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:20px}section{background:#fff;border:1px solid var(--line);border-radius:12px;padding:23px;box-shadow:0 2px 8px #173b5c10}h2{color:var(--navy);margin:0 0 8px;font-size:20px}.help{margin:0 0 19px;color:#526577;line-height:1.45}label{display:block;font-weight:650;margin:15px 0 6px}input,textarea,select{width:100%;border:1px solid #aebfce;border-radius:7px;padding:10px;font:inherit;background:white}textarea{min-height:82px;resize:vertical}button,.button{display:inline-block;border:0;background:#176b42;color:#fff;padding:11px 16px;border-radius:7px;font-weight:700;font:inherit;text-decoration:none;cursor:pointer;margin-top:20px}.button.secondary{background:var(--navy);margin:0 8px 0 0}.result{margin-top:25px}.result h2{margin-bottom:11px}.result li{margin:10px 0}.small{font-size:14px;color:#627384}.table-wrap{overflow:auto;border:1px solid var(--line);border-radius:8px;background:#fff}table{border-collapse:collapse;width:100%;font-size:14px}th,td{border-bottom:1px solid #dde6ee;padding:8px 10px;text-align:left;vertical-align:top;white-space:nowrap}th{background:var(--blue);position:sticky;top:0}@media(max-width:600px){main{padding:0 14px}.grid{grid-template-columns:1fr}}</style></head><body><header><h1>Nava Vaibhva Bill Book</h1><p>Local report generator — nothing is sent outside this computer.</p></header><main>
{% if message %}<div class="notice">{{ message }}</div>{% endif %}{% if error %}<div class="notice error">{{ error }}</div>{% endif %}<div class="grid"><section><h2>Monthly bill report</h2><p class="help">Select one or more monthly bank statements. A complete monthly workbook is created for every registered flat.</p><form method="post" action="{{ url_for('monthly') }}" enctype="multipart/form-data"><label>Bank statement file(s)</label><input required name="statements" type="file" multiple accept=".xls,.xlsx,.xlsm,.csv"><label>Water bill per flat (optional)</label><input name="water_bill" type="number" min="0" step="0.01" placeholder="For example, 287"><p class="small">Leave blank to use the configured value.</p><button>Generate monthly report</button></form></section><section><h2>Collection history</h2><p class="help">Select statements covering the period, then optionally enter one or more flats. Leave flats blank to create history for every flat.</p><form method="post" action="{{ url_for('history') }}" enctype="multipart/form-data"><label>Bank statement file(s)</label><input required name="statements" type="file" multiple accept=".xls,.xlsx,.xlsm,.csv"><label>Flat number(s) (optional)</label><textarea name="flats" placeholder="A001, A002, B303&#10;or paste one flat per line"></textarea><p class="small">Separate flat numbers with commas, spaces, or new lines.</p><button>Generate collection history</button></form></section></div>
{% if reports %}<section class="result"><h2>Your report{{ 's' if reports|length > 1 else '' }} are ready</h2><ul>{% for report in reports %}<li><strong>{{ report.name }}</strong><br><a class="button secondary" href="{{ url_for('preview', filename=report.name) }}">View first</a><a class="button" href="{{ url_for('download', filename=report.name) }}">Download Excel</a></li>{% endfor %}</ul></section>{% endif %}{% if preview %}<section class="result"><p><a class="button secondary" href="{{ url_for('index') }}">Back to reports</a><a class="button" href="{{ url_for('download', filename=preview.filename) }}">Download Excel</a></p><h2>{{ preview.filename }}</h2><form method="get"><label>Worksheet</label><select name="sheet" onchange="this.form.submit()">{% for name in preview.sheets %}<option value="{{ name }}" {% if name == preview.selected %}selected{% endif %}>{{ name }}</option>{% endfor %}</select></form><p class="small">Showing the first 100 rows and 25 columns. Download the workbook for the full report.</p><div class="table-wrap"><table><tr>{% for cell in preview.header %}<th>{{ cell }}</th>{% endfor %}</tr>{% for row in preview.rows %}<tr>{% for cell in row %}<td>{{ cell }}</td>{% endfor %}</tr>{% endfor %}</table></div></section>{% endif %}</main></body></html>"""


def create_app(project_root: Path | None = None) -> Flask:
    root = (project_root or Path.cwd()).resolve()
    app = Flask(__name__)

    def config_and_registry():
        config_path = root / "config.yaml"
        if not config_path.exists():
            raise ValueError("config.yaml is missing. Copy config.example.yaml to config.yaml first.")
        config = Config.load(config_path)
        if not config.output_dir.is_absolute(): config.output_dir = root / config.output_dir
        if not config.input_dir.is_absolute(): config.input_dir = root / config.input_dir
        return config, _load_flats_registry(config)

    def save_uploads(files, folder: Path) -> list[Path]:
        saved = []
        for item in files:
            if not item or not item.filename: continue
            name = Path(item.filename).name
            if Path(name).suffix.lower() not in ALLOWED_EXTENSIONS: raise ValueError(f"{name}: please select an Excel or CSV statement file.")
            destination = folder / name
            item.save(destination)
            saved.append(destination)
        if not saved: raise ValueError("Select at least one bank statement file.")
        return saved

    @app.get("/")
    def index(): return render_template_string(PAGE)

    @app.post("/monthly")
    def monthly():
        try:
            config, registry = config_and_registry()
            raw = request.form.get("water_bill", "").strip()
            water_bill = float(raw) if raw else None
            if water_bill is not None and water_bill < 0: raise ValueError("Water bill cannot be negative.")
            with tempfile.TemporaryDirectory(prefix="nv-billbook-") as folder:
                reports = [_process_statement(path, config, None, registry, water_bill) for path in save_uploads(request.files.getlist("statements"), Path(folder))]
            reports = [path for path in reports if path]
            if not reports: raise ValueError("No monthly reports were created from the selected statements.")
            return render_template_string(PAGE, reports=reports, message="Monthly report generation is complete.")
        except Exception as exc: return render_template_string(PAGE, error=str(exc)), 400

    @app.post("/history")
    def history():
        try:
            config, registry = config_and_registry()
            if not registry: raise ValueError("A flats registry is required to create collection history.")
            with tempfile.TemporaryDirectory(prefix="nv-history-") as folder:
                statements = save_uploads(request.files.getlist("statements"), Path(folder))
                transactions = _load_statements(statements[0] if len(statements) == 1 else Path(folder), config)
                months = month_keys_from_transactions(transactions)
                if not months: raise ValueError("No transaction months were found in the selected statements.")
                water_bills = registry.meta.get("water_bills_by_month")
                if not isinstance(water_bills, Mapping): raise ValueError("Add monthly water bills to flats.yaml before creating collection history.")
                flat_text = request.form.get("flats", "").replace("\n", ",").replace(" ", ",")
                selection = _parse_flat_filter(flat_text, registry) if flat_text.strip(",") else None
                credits = split_by_type(classify_transactions(transactions, config, registry))["credits"]
                transaction_history = build_flat_transaction_history(credits, registry)
                reconciliation = build_monthly_reconciliation(credits, registry, months, water_bills)
                if selection:
                    transaction_history = transaction_history[transaction_history["Flat"].isin(selection)]
                    reconciliation = reconciliation[reconciliation["Flat"].isin(selection)]
                suffix = selection[0] if selection and len(selection) == 1 else "selected" if selection else "collection"
                report = write_collection_history_report(config.output_dir / f"{months[0]}_to_{months[-1]}_{suffix}_history.xlsx", transaction_history, reconciliation, build_period_summary(reconciliation))
            return render_template_string(PAGE, reports=[report], message="Collection history is ready.")
        except Exception as exc: return render_template_string(PAGE, error=str(exc)), 400

    def report_path(filename: str) -> Path:
        config, _ = config_and_registry()
        candidate = (config.output_dir / Path(filename).name).resolve()
        if candidate.parent != config.output_dir.resolve() or not candidate.is_file(): abort(404)
        return candidate

    @app.get("/reports/<filename>/download")
    def download(filename: str):
        path = report_path(filename)
        return send_from_directory(path.parent, path.name, as_attachment=True)

    @app.get("/reports/<filename>/preview")
    def preview(filename: str):
        path = report_path(filename)
        workbook = load_workbook(path, data_only=True, read_only=True)
        selected = request.args.get("sheet") or workbook.sheetnames[0]
        if selected not in workbook.sheetnames: selected = workbook.sheetnames[0]
        sheet = workbook[selected]
        table = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(max_row=100, max_col=25, values_only=True)]
        header, rows = (table[0], table[1:]) if table else ([], [])
        sheets = workbook.sheetnames
        workbook.close()
        return render_template_string(PAGE, preview={"filename": path.name, "sheets": sheets, "selected": selected, "header": header, "rows": rows})

    return app


def main() -> None:
    app = create_app()
    print("Nava Vaibhva Bill Book is ready at http://127.0.0.1:5000")
    app.run(host="127.0.0.1", port=5000, debug=False)


if __name__ == "__main__": main()
