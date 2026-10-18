"""Build vendor-wise payment history workbooks."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter


TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
ALT_ROW_FILL = PatternFill("solid", fgColor="F7F9FB")
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
HEADER_FONT = Font(bold=True)
THIN_GREY = Side(style="thin", color="D9E1F2")
CELL_BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)


def normalize_vendor(value: object) -> str:
    """Create a stable grouping key while retaining readable display names."""
    text = "" if value is None else str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text.casefold()


def _display_vendor(value: object) -> str:
    text = "" if value is None else str(value).strip()
    return re.sub(r"\s+", " ", text) or "Unassigned / Other"


def _prepare_transactions(debits: pd.DataFrame) -> pd.DataFrame:
    if debits.empty:
        return pd.DataFrame(
            columns=[
                "Date",
                "Vendor / Payee",
                "Purpose / Type",
                "Category",
                "Reference",
                "Debit (₹)",
                "Balance (₹)",
                "Source File",
            ]
        )

    result = pd.DataFrame(
        {
            "Date": debits["date"].map(
                lambda value: value.strftime("%d/%m/%Y") if hasattr(value, "strftime") else ""
            ),
            "Vendor / Payee": debits["counterparty"].map(_display_vendor),
            "Purpose / Type": debits["purpose"].fillna(""),
            "Category": debits["category"].fillna(""),
            "Reference": debits["reference"].fillna(""),
            "Debit (₹)": debits["debit"],
            "Balance (₹)": debits["balance"],
            "Source File": debits["source_file"].fillna(""),
        }
    )
    return result


def build_vendor_summary(debits: pd.DataFrame) -> pd.DataFrame:
    """Return one row per normalized vendor/payee."""
    if debits.empty:
        return pd.DataFrame(
            columns=[
                "Vendor / Payee",
                "Total Paid (₹)",
                "Payments",
                "First Payment",
                "Latest Payment",
                "Categories",
            ]
        )

    rows: list[dict[str, object]] = []
    working = debits.copy()
    working["_vendor_key"] = working["counterparty"].map(normalize_vendor)
    for _, group in working.groupby("_vendor_key", sort=False, dropna=False):
        dates = group["date"].dropna()
        display_names = [_display_vendor(value) for value in group["counterparty"]]
        vendor_name = next((name for name in display_names if name != "Unassigned / Other"), "Unassigned / Other")
        categories = sorted({str(value) for value in group["category"] if pd.notna(value) and str(value).strip()})
        rows.append(
            {
                "Vendor / Payee": vendor_name,
                "Total Paid (₹)": group["debit"].sum(),
                "Payments": len(group),
                "First Payment": min(dates).strftime("%d/%m/%Y") if not dates.empty else "",
                "Latest Payment": max(dates).strftime("%d/%m/%Y") if not dates.empty else "",
                "Categories": ", ".join(categories),
            }
        )
    return pd.DataFrame(rows).sort_values("Total Paid (₹)", ascending=False, ignore_index=True)


def filter_vendors(debits: pd.DataFrame, vendor_filter: str | None) -> pd.DataFrame:
    if not vendor_filter or debits.empty:
        return debits
    requested = [normalize_vendor(item) for item in vendor_filter.split(",") if item.strip()]
    keys = debits["counterparty"].map(normalize_vendor)
    mask = keys.map(lambda key: any(requested_name in key for requested_name in requested))
    return debits.loc[mask].copy()


def _style_table(worksheet, frame: pd.DataFrame, header_row: int = 3) -> None:
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(max(1, len(frame.columns)))}{worksheet.max_row}"
    for cell in worksheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        if (row_number - header_row) % 2 == 0:
            for cell in worksheet[row_number]:
                cell.fill = ALT_ROW_FILL
        for cell in worksheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 3, 4, 6, 8})
            header = worksheet.cell(row=header_row, column=cell.column).value
            if header in {"Total Paid (₹)", "Debit (₹)", "Balance (₹)"} and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
    for index, column in enumerate(frame.columns, start=1):
        values = [str(column)] + [str(value) for value in frame[column].tolist() if pd.notna(value)]
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(len(value) for value in values) + 2, 45)


def write_vendor_report(
    output_path: Path,
    debits: pd.DataFrame,
    period_label: str,
    vendor_filter: str | None = None,
) -> Path:
    """Write a vendor summary and detailed vendor transaction history."""
    filtered = filter_vendors(debits, vendor_filter)
    summary = build_vendor_summary(filtered)
    transactions = _prepare_transactions(filtered)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Vendor Summary", startrow=2, index=False)
        transactions.to_excel(writer, sheet_name="Vendor Transactions", startrow=2, index=False)
        for sheet_name, frame, title in (
            ("Vendor Summary", summary, "Vendor Payment Summary"),
            ("Vendor Transactions", transactions, "Vendor Payment Transactions"),
        ):
            worksheet = writer.sheets[sheet_name]
            worksheet["A1"] = f"{title} — {period_label}"
            worksheet["A1"].font = TITLE_FONT
            worksheet["A1"].fill = TITLE_FILL
            worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(frame.columns)))
            _style_table(worksheet, frame)
    return output_path
