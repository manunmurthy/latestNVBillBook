"""Tests for multi-month flat collection history calculations."""

from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from nv_billbook.collection_history import (
    MissingWaterBillError,
    build_flat_transaction_history,
    build_monthly_reconciliation,
    build_period_summary,
    month_keys_from_transactions,
    month_keys_for_period,
)
from nv_billbook.flats_registry import FlatInfo, FlatsRegistry
from nv_billbook.history_reporter import write_collection_history_report


@pytest.fixture
def registry() -> FlatsRegistry:
    return FlatsRegistry(
        meta={"maintenance_rate_per_sqft": 2.5},
        flats={
            "A001": FlatInfo("A001", "A", 1000),
            "A002": FlatInfo("A002", "A", 1000),
        },
        unmapped_payers=[],
    )


@pytest.fixture
def credits() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"date": date(2026, 3, 5), "flat_no": "A001", "counterparty": "Payer A", "category": "Maintenance", "credit": 2787},
            {"date": date(2026, 4, 5), "flat_no": "A001", "counterparty": "Payer A", "category": "Maintenance", "credit": 2500},
            {"date": date(2026, 5, 5), "flat_no": "A001", "counterparty": "Payer A", "category": "Maintenance", "credit": 3100},
            {"date": date(2026, 3, 6), "flat_no": None, "counterparty": "Unknown", "category": "Maintenance", "credit": 9999},
        ]
    )


def test_history_and_reconciliation_include_every_flat(
    registry: FlatsRegistry, credits: pd.DataFrame
) -> None:
    months = month_keys_for_period(date(2026, 3, 1), date(2026, 5, 31))
    water_bills = {"2026-03": 287, "2026-04": 287, "2026-05": 287}

    history = build_flat_transaction_history(credits, registry)
    monthly = build_monthly_reconciliation(credits, registry, months, water_bills)
    summary = build_period_summary(monthly)

    assert len(history) == 3
    assert list(history["Flat"]) == ["A001", "A001", "A001"]
    assert len(monthly) == 6
    assert list(monthly[monthly["Flat"] == "A001"]["Status"]) == ["Paid", "Short", "Excess"]
    assert list(monthly[monthly["Flat"] == "A002"]["Status"]) == ["Pending"] * 3
    assert list(summary["Flat"]) == ["A001", "A002"]
    assert list(summary["Overall Status"]) == ["Excess", "Pending"]
    assert summary.loc[0, "Total Extra Paid"] == 26.0


def test_months_come_from_transaction_dates_not_statement_boundaries() -> None:
    transactions = pd.DataFrame(
        {"date": [date(2026, 3, 1), date(2026, 4, 15), date(2026, 5, 31)]}
    )

    assert month_keys_from_transactions(transactions) == ["2026-03", "2026-04", "2026-05"]


def test_reconciliation_rejects_missing_monthly_water_bill(
    registry: FlatsRegistry, credits: pd.DataFrame
) -> None:
    with pytest.raises(MissingWaterBillError, match="2026-04"):
        build_monthly_reconciliation(
            credits,
            registry,
            ["2026-03", "2026-04"],
            {"2026-03": 287},
        )


def test_history_workbook_has_one_sheet_per_flat(
    tmp_path, registry: FlatsRegistry, credits: pd.DataFrame
) -> None:
    months = ["2026-03", "2026-04", "2026-05"]
    monthly = build_monthly_reconciliation(
        credits, registry, months, {month: 287 for month in months}
    )
    output_path = write_collection_history_report(
        tmp_path / "history.xlsx",
        build_flat_transaction_history(credits, registry),
        monthly,
        build_period_summary(monthly),
    )

    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == ["A001", "A002"]
    assert workbook["A001"]["A1"].value == "Flat A001 — Collection History"
    assert workbook["A002"]["A1"].value == "Flat A002 — Collection History"


def test_history_workbook_can_be_limited_to_one_flat(
    tmp_path, registry: FlatsRegistry, credits: pd.DataFrame
) -> None:
    months = ["2026-03", "2026-04", "2026-05"]
    monthly = build_monthly_reconciliation(
        credits, registry, months, {month: 287 for month in months}
    )
    flat_monthly = monthly[monthly["Flat"] == "A001"]
    flat_history = build_flat_transaction_history(credits, registry)
    flat_history = flat_history[flat_history["Flat"] == "A001"]
    flat_summary = build_period_summary(flat_monthly)

    output_path = write_collection_history_report(
        tmp_path / "history_a001.xlsx",
        flat_history,
        flat_monthly,
        flat_summary,
    )

    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == ["A001"]
    assert workbook["A001"]["A1"].value == "Flat A001 — Collection History"
    assert workbook["A001"]["A4"].value == "A001"
